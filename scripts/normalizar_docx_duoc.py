from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import sys

BLUE="17365D"
LIGHT="DCE6F1"
GRAY="5F6B7A"
CONTROL="Control documental 03-10-2026 · Scrum/Trello vigente · PO Alexander Patiño · SM Samuel Solís · Dev Camilo González"

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr()
    shd=OxmlElement("w:shd")
    shd.set(qn("w:fill"),fill)
    tcPr.append(shd)

def normalize(path):
    doc=Document(path)
    for sec in doc.sections:
        sec.top_margin=Inches(0.7)
        sec.bottom_margin=Inches(0.65)
        sec.left_margin=Inches(0.8)
        sec.right_margin=Inches(0.8)

    styles=doc.styles
    if "Normal" in styles:
        styles["Normal"].font.name="Arial"
        styles["Normal"].font.size=Pt(10.5)

    for name,size in [("Title",24),("Heading 1",16),("Heading 2",13),("Heading 3",11)]:
        if name in styles:
            styles[name].font.name="Arial"
            styles[name].font.size=Pt(size)
            styles[name].font.bold=True
            styles[name].font.color.rgb=RGBColor.from_string(BLUE)

    for tbl in doc.tables:
        try:
            tbl.style="Table Grid"
        except Exception:
            pass
        if tbl.rows:
            for c in tbl.rows[0].cells:
                shade(c,LIGHT)
                for p in c.paragraphs:
                    for r in p.runs:
                        r.bold=True
                        r.font.name="Arial"
                        r.font.color.rgb=RGBColor.from_string(BLUE)

    # Preserva logos/cabeceras existentes. Solo agrega control documental al pie.
    for sec in doc.sections:
        footer=sec.footer
        existing=" ".join(p.text for p in footer.paragraphs)
        if CONTROL not in existing:
            p=footer.add_paragraph()
            p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
            r=p.add_run(CONTROL)
            r.font.name="Arial"
            r.font.size=Pt(7)
            r.font.color.rgb=RGBColor.from_string(GRAY)

    doc.save(path)

if __name__=="__main__":
    root=Path(sys.argv[1] if len(sys.argv)>1 else ".")
    files=[
        p for p in root.rglob("*.docx")
        if ".git" not in p.parts and "node_modules" not in p.parts
    ]
    ok=0
    for f in files:
        try:
            normalize(f)
            print("OK",f)
            ok+=1
        except Exception as e:
            print("ERROR",f,e)
    print("Procesados:",ok,"de",len(files))
